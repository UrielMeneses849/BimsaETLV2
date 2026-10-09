import tempfile
import unittest
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU, points_to_pixels

from BIMSA_pipeline import process_bimsa


def _record(origin=None, clave="OC000001", precision="Calle"):
    origin_xml = f"<Origen>{origin}</Origen>" if origin is not None else ""
    return f"""
    <datos>
      {origin_xml}
      <Clave_Proyecto>{clave}</Clave_Proyecto>
      <Fecha_Publicacion>01/10/2026</Fecha_Publicacion>
      <Tipo_Proyecto>Proyecto contratado</Tipo_Proyecto>
      <Proyecto>Proyecto de prueba</Proyecto>
      <Region>Centro</Region>
      <proy_ubicacionlatitud>19.432608</proy_ubicacionlatitud>
      <proy_ubicacionlongitud>-99.133209</proy_ubicacionlongitud>
      <Precision>{precision}</Precision>
      <Compania>Compañía de prueba</Compania>
    </datos>
    """


def _workbook_for(*records):
    xml = f"<bobras>{''.join(records)}</bobras>"
    with tempfile.TemporaryDirectory() as output_root:
        _, content = process_bimsa(
            xml,
            "CLASICO",
            output_root=output_root,
            guardar_json=False,
        )
    return load_workbook(BytesIO(content), data_only=True)


class ClassicoSourceSheetsTest(unittest.TestCase):
    def test_mixed_sources_create_two_specialized_sheets(self):
        workbook = _workbook_for(
            _record("Construleads", "OC000001"),
            _record(" Explorer ", "EX000001", "Predio"),
        )

        self.assertEqual(workbook.sheetnames, ["Construleads", "Explorer"])

        construleads = workbook["Construleads"]
        explorer = workbook["Explorer"]
        construleads_headers = [
            construleads.cell(3, column).value
            for column in range(1, construleads.max_column + 1)
        ]
        explorer_headers = [
            explorer.cell(3, column).value
            for column in range(1, explorer.max_column + 1)
        ]

        for header in ("Latitud", "Longitud", "Precisión"):
            self.assertIn(header, construleads_headers)
            self.assertIn(header, explorer_headers)
        self.assertEqual(
            construleads_headers[-3:],
            ["Latitud", "Longitud", "Precisión"],
        )
        self.assertEqual(
            explorer_headers[-3:],
            ["Latitud", "Longitud", "Precisión"],
        )

        self.assertIn("Compania", construleads_headers)
        self.assertNotIn("Compania", explorer_headers)
        self.assertNotIn("Clave Compania", explorer_headers)

        latitude_column = explorer_headers.index("Latitud") + 1
        longitude_column = explorer_headers.index("Longitud") + 1
        precision_column = explorer_headers.index("Precisión") + 1
        self.assertEqual(explorer.cell(4, latitude_column).value, 19.432608)
        self.assertEqual(explorer.cell(4, longitude_column).value, -99.133209)
        self.assertEqual(explorer.cell(4, precision_column).value, "Predio")

        groups = [
            explorer.cell(2, column)
            for column in range(1, explorer.max_column + 1)
            if explorer.cell(2, column).value
        ]
        self.assertEqual(
            [cell.value for cell in groups],
            [
                "Información General",
                "Ubicación del Proyecto",
                "Caracteristicas del Proyecto",
                "Participantes y Contactos",
                "Detalles Adicionales",
                "Geolocalización",
            ],
        )
        self.assertTrue(groups[3].fill.fgColor.rgb.endswith("E7E6E6"))
        self.assertTrue(groups[4].fill.fgColor.rgb.endswith("ED7D31"))
        self.assertTrue(groups[5].fill.fgColor.rgb.endswith("E7E6E6"))

        self.assertEqual(len(construleads._images), 1)
        self.assertEqual(len(explorer._images), 1)
        for worksheet in (construleads, explorer):
            image = worksheet._images[0]
            visible_height_px = image.anchor.ext.cy / pixels_to_EMU(1)
            expected_offset = pixels_to_EMU(
                (points_to_pixels(worksheet.row_dimensions[1].height) - visible_height_px) / 2
            )
            self.assertAlmostEqual(
                image.anchor._from.rowOff,
                expected_offset,
                delta=1,
            )
            visible_width_px = image.anchor.ext.cx / pixels_to_EMU(1)
            expected_horizontal_offset = pixels_to_EMU(
                ((worksheet.column_dimensions["A"].width * 8) - visible_width_px) / 2
            )
            self.assertAlmostEqual(
                image.anchor._from.colOff,
                expected_horizontal_offset,
                delta=1,
            )

        construleads_widths = {
            header: construleads.column_dimensions[
                get_column_letter(column)
            ].width
            for column, header in enumerate(construleads_headers, start=1)
        }
        explorer_widths = {
            header: explorer.column_dimensions[get_column_letter(column)].width
            for column, header in enumerate(explorer_headers, start=1)
        }
        self.assertEqual(construleads_widths["Inversión (Pesos Mexicanos)"], 31)
        self.assertEqual(construleads_widths["Sup Urbanizada (m²)"], 13)
        self.assertEqual(construleads_widths["Descripcion"], 127)
        self.assertAlmostEqual(explorer_widths["Titulo 1"], 13.43, places=2)
        self.assertAlmostEqual(explorer_widths["Acabados"], 122.43, places=2)
        self.assertEqual(construleads_widths["Latitud"], 16)
        self.assertEqual(construleads_widths["Longitud"], 16)
        self.assertEqual(construleads_widths["Precisión"], 18)

    def test_only_explorer_omits_construleads_sheet(self):
        workbook = _workbook_for(_record("Explorer", "EX000001"))
        self.assertEqual(workbook.sheetnames, ["Explorer"])

    def test_only_construleads_omits_explorer_sheet(self):
        workbook = _workbook_for(_record("Construleads", "OC000001"))
        self.assertEqual(workbook.sheetnames, ["Construleads"])

    def test_missing_origin_remains_backward_compatible(self):
        workbook = _workbook_for(_record(None, "OC000001"))
        self.assertEqual(workbook.sheetnames, ["Construleads"])


if __name__ == "__main__":
    unittest.main()
